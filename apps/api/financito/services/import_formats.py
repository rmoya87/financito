from __future__ import annotations
import csv,re
from io import BytesIO,StringIO
from pathlib import Path
from defusedxml import ElementTree as ET
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from .imports import import_csv

def _csv_bytes(rows:list[dict])->bytes:
    if not rows:return b""
    out=StringIO();w=csv.DictWriter(out,fieldnames=["Fecha","Concepto","Importe","Moneda","Comercio"],delimiter=";");w.writeheader()
    for r in rows:w.writerow(r)
    return out.getvalue().encode()

def _xlsx(content:bytes)->bytes:
    wb=load_workbook(BytesIO(content),read_only=True,data_only=True);ws=wb.active
    values=ws.iter_rows(values_only=True);headers=[str(x or "").strip() for x in next(values)]
    rows=[]
    for row in values:
        d={headers[i]:row[i] for i in range(min(len(headers),len(row)))}
        def pick(*names):
            for n in names:
                for k,v in d.items():
                    if k.lower().replace(" ","") == n and v not in (None,""):return v
            return ""
        rows.append({"Fecha":pick("fecha","date","bookingdate"),"Concepto":pick("concepto","descripcion","description","detalle"),"Importe":pick("importe","amount","cantidad"),"Moneda":pick("moneda","currency") or "EUR","Comercio":pick("comercio","merchant","beneficiario")})
    return _csv_bytes(rows)

def _qif(text:str)->bytes:
    rows=[];cur={}
    for line in text.splitlines():
        if not line:continue
        code,val=line[:1],line[1:].strip()
        if code=="D":cur["Fecha"]=val
        elif code=="T":cur["Importe"]=val
        elif code=="P":cur["Comercio"]=val
        elif code=="M":cur["Concepto"]=val
        elif code=="^":
            if cur:cur.setdefault("Concepto",cur.get("Comercio","Movimiento"));cur.setdefault("Moneda","EUR");rows.append(cur)
            cur={}
    return _csv_bytes(rows)

def _ofx(text:str)->bytes:
    rows=[]
    for block in re.findall(r"<STMTTRN>(.*?)(?:</STMTTRN>|<STMTTRN>)",text,re.S|re.I):
        def tag(name):
            m=re.search(fr"<{name}>([^<\r\n]+)",block,re.I);return m.group(1).strip() if m else ""
        dt=tag("DTPOSTED")[:8]
        if len(dt)==8:dt=f"{dt[6:8]}/{dt[4:6]}/{dt[:4]}"
        rows.append({"Fecha":dt,"Importe":tag("TRNAMT"),"Concepto":tag("MEMO") or tag("NAME") or "Movimiento","Comercio":tag("NAME"),"Moneda":"EUR"})
    return _csv_bytes(rows)

def _local(el,tag):return [x for x in el.iter() if x.tag.rsplit("}",1)[-1]==tag]
def _camt(content:bytes)->bytes:
    root=ET.fromstring(content);rows=[]
    for entry in _local(root,"Ntry"):
        amts=_local(entry,"Amt");dates=_local(entry,"Dt");infos=_local(entry,"AddtlNtryInf");cd=_local(entry,"CdtDbtInd")
        if not amts or not dates:continue
        amount=amts[0].text or "0"
        if cd and (cd[0].text or "").upper()=="DBIT":amount="-"+amount.lstrip("-")
        rows.append({"Fecha":dates[0].text or "","Importe":amount,"Concepto":infos[0].text if infos else "Movimiento","Comercio":"","Moneda":amts[0].attrib.get("Ccy","EUR")})
    return _csv_bytes(rows)

def _mt940(text:str)->bytes:
    rows=[];current=None
    for line in text.splitlines():
        if line.startswith(":61:"):
            m=re.match(r":61:(\d{6})(?:\d{4})?([CD])([\d,\.]+)",line)
            if m:
                yy=int(m.group(1)[:2]);year=2000+yy if yy<70 else 1900+yy;date=f"{m.group(1)[4:6]}/{m.group(1)[2:4]}/{year}"
                amount=m.group(3).replace(",",".");amount=("-" if m.group(2)=="D" else "")+amount
                current={"Fecha":date,"Importe":amount,"Concepto":"Movimiento","Comercio":"","Moneda":"EUR"};rows.append(current)
        elif line.startswith(":86:") and current:current["Concepto"]=line[4:].strip()
    return _csv_bytes(rows)

def import_statement(session:Session,account_id:str,filename:str,content:bytes):
    ext=Path(filename).suffix.lower()
    if ext==".csv":normalized=content
    elif ext in {".xlsx",".xlsm"}:normalized=_xlsx(content)
    elif ext==".qif":normalized=_qif(content.decode("utf-8",errors="replace"))
    elif ext==".ofx":normalized=_ofx(content.decode("utf-8",errors="replace"))
    elif ext in {".xml",".camt",".053"}:normalized=_camt(content)
    elif ext in {".mt940",".sta"}:normalized=_mt940(content.decode("utf-8",errors="replace"))
    else:raise ValueError(f"Unsupported statement format: {ext}")
    return import_csv(session,account_id,normalized,filename)
