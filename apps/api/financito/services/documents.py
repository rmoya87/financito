from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
import json
import mimetypes
import re

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader
from PIL import Image
from pillow_heif import register_heif_opener
import pytesseract
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import settings
from ..models import ActionItem, Document, ExtractedFact

register_heif_opener()


@dataclass(frozen=True)
class IndexedDocument:
    document: Document
    facts_created: int
    chunks_created: int


def safe_path(path: Path) -> Path:
    expanded = path.expanduser()
    if expanded.is_symlink():
        raise ValueError("Symlink documents are not accepted")
    resolved = expanded.resolve(strict=True)
    vault = settings.vault_dir.resolve()
    if resolved != vault and vault not in resolved.parents:
        raise ValueError("Document is outside configured vault")
    return resolved


def extract_content(path: Path) -> tuple[str, int, list[str] | None]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n\n".join(pages), len(pages), pages
    if suffix in {".txt", ".csv", ".json"}:
        return path.read_text(encoding="utf-8", errors="replace"), 1, None
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        return "\n".join(p.text for p in doc.paragraphs), 1, None
    if suffix in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        lines: list[str] = []
        for ws in wb.worksheets:
            lines.append(f"[{ws.title}]")
            for row in ws.iter_rows(values_only=True):
                lines.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(lines), 1, None
    if suffix in {".png", ".jpg", ".jpeg", ".heic", ".tiff", ".bmp"}:
        image = Image.open(path)
        return pytesseract.image_to_string(image, lang="spa+eng"), 1, None
    raise ValueError(f"Unsupported document type: {suffix}")


def extract_contract_facts(text: str) -> list[dict]:
    facts: list[dict] = []
    patterns = [
        ("cancellation_notice_days", r"(?:preaviso|antelaci[oó]n)\D{0,50}(\d{1,3})\s*d[ií]as", "days"),
        ("early_exit_penalty", r"(?:penalizaci[oó]n|comisi[oó]n)\D{0,80}(\d+[\.,]?\d*)\s*(?:€|euros?)", "EUR"),
        ("annual_cost", r"(?:prima anual|coste anual|cuota anual)\D{0,60}(\d+[\.,]?\d*)\s*(?:€|euros?)", "EUR"),
    ]
    lowered = text.lower()
    for key, pattern, unit in patterns:
        match = re.search(pattern, lowered, re.IGNORECASE)
        if match:
            raw = match.group(1).replace(",", ".")
            facts.append({"fact_type": "contract_term", "key": key, "value": raw, "unit": unit, "confidence": 0.72})
    for key, pattern in [
        ("permanence_end_date", r"permanencia.{0,80}(\d{1,2}/\d{1,2}/\d{4})"),
        ("renewal_date", r"renovaci[oó]n.{0,80}(\d{1,2}/\d{1,2}/\d{4})"),
    ]:
        match = re.search(pattern, lowered, re.IGNORECASE)
        if match:
            facts.append({"fact_type": "contract_term", "key": key, "value": match.group(1), "unit": "date", "confidence": 0.68})
    return facts


def index_document(session: Session, source_path: str, document_type: str = "unknown") -> IndexedDocument:
    path = safe_path(Path(source_path))
    digest = sha256(path.read_bytes()).hexdigest()
    existing = session.scalar(select(Document).where(Document.sha256 == digest))
    if existing:
        from .rag import index_document_chunks

        chunks = index_document_chunks(session, existing)
        return IndexedDocument(existing, 0, chunks)

    extracted_text, page_count, pages = extract_content(path)
    doc = Document(
        file_path=str(path),
        file_name=path.name,
        mime_type=mimetypes.guess_type(path.name)[0],
        sha256=digest,
        document_type=document_type,
        status="indexed",
        page_count=page_count,
        extracted_text=extracted_text,
    )
    session.add(doc)
    session.flush()

    count = 0
    for fact in extract_contract_facts(extracted_text):
        session.add(
            ExtractedFact(
                document_id=doc.id,
                fact_type=fact["fact_type"],
                key=fact["key"],
                value_json=json.dumps({"value": fact["value"], "unit": fact["unit"]}, ensure_ascii=False),
                confidence=str(fact["confidence"]),
                status="inferred",
                user_verified=False,
            )
        )
        count += 1

    if count:
        session.add(
            ActionItem(
                action_type="review_document_evidence",
                title=f"Revisar {count} dato(s) contractual(es) extraído(s) de {path.name}",
                related_entity_type="document",
                related_entity_id=doc.id,
                priority="high",
                source_type="document",
                source_ref=doc.id,
                notes="Los datos extraídos son inferidos y no deben usarse como evidencia confirmada hasta su revisión.",
            )
        )

    from .rag import index_document_chunks

    chunks = index_document_chunks(session, doc, pages)
    return IndexedDocument(doc, count, chunks)
